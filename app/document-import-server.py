"""Local-only document extraction for Anju. Original files stay in the browser."""
import csv
import io
import json
import re
import sys
import threading
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

sys.path.insert(0, str(Path(__file__).parent / 'document-deps'))
ALIASES = {
    'name': ['项目名称', '商品名称', '商品全名', '产品名称', '柜体名称', '报价项目', '货物名称', '品名', '存货名称', '材料名称', '项目', '名称', '柜子'],
    'unit': ['计价单位', '单位'], 'qty': ['工程数量', '工程量', '计价量', '面积', '平方', '数量'],
    'price': ['客户单价', '报价单价', '销售单价', '单价'],
    'amount': ['总金额', '金额', '合价', '小计', '总价'],
    'spec': ['规格型号', '型号规格', '规格', '尺寸'], 'model': ['型号编码', '商品编码', '花色编码', '货号', '编码', '型号', '花色'],
    'room': ['所属房间', '空间区域', '所属', '空间', '房间', '位置'], 'remark': ['备注', '说明', '要求'],
    'serial': ['序号', '行号'],
    'category': ['项目分类', '大类', '分类', '类别'], 'brand': ['品牌'],
    'substrate': ['板材类型', '基材', '材质'], 'thickness': ['厚度'], 'color': ['颜色', '色号', '花色颜色'],
}

def norm(value):
    return re.sub(r'[\s（）()：:/\\￥¥元㎡m²*\d._-]', '', str(value or '')).lower()

def field(value):
    s = norm(value)
    for key, names in ALIASES.items():
        if s and any(s == norm(n) for n in names):
            return key
    if s.startswith('规格材质'):
        return 'spec'
    return None

def number(value):
    if isinstance(value, (int, float)):
        return float(value)
    s = re.sub(r'\s+', '', str(value or '')).replace(',', '').replace('，', '')
    m = re.fullmatch(r'[￥¥]?\s*(-?\d+(?:\.\d+)?)\s*(?:元|㎡|m²|平方米|平方|张|个|套|米|m|kg)?', s, re.I)
    return float(m[1]) if m else None

def infer_material(row):
    s = ' '.join(str(row.get(k) or '') for k in ['name', 'spec', 'model', 'remark'])
    if not row.get('substrate'):
        row['substrate'] = next((x for x in ['多层板', '颗粒板', '生态板', '欧松板', '密度板', '实木板', '禾香板', 'F-OSB'] if x in s), '')
    if not row.get('thickness'):
        m = re.search(r'(?:厚度\s*[:：]?\s*)?(\d{1,2}(?:\.\d+)?)\s*(?:mm|厚)', s, re.I)
        row['thickness'] = m[1] if m else ''
    if not row.get('brand'):
        row['brand'] = next((x for x in ['兔宝宝', '千年舟', '大亚圣象', '大亚', '鲁丽', '福人', '禾香', '诺米', '百隆', '海蒂诗', '东泰', '悍高', '特怀克', '欣锐嘉'] if x in s), '')
    if not row.get('color'):
        code = re.search(r'([A-Z]{1,3}\d{3,}(?:[- ]?[A-Z]{2,8})?)', s)
        row['color'] = code[1] if code else ''
    return row

def table_rows(rows):
    header, start = {}, -1
    for i, r in enumerate(rows[:100]):
        mapping = {field(v): c for c, v in enumerate(r) if field(v)}
        if 'name' in mapping and len(mapping) >= 3 and len(mapping) > len(header):
            header, start = mapping, i
    if start < 0:
        return [], '未定位到项目表头，请在预览中指定列对应关系。'
    result, room, category = [], '', ''
    for r in rows[start + 1:]:
        def get(key):
            c = header.get(key)
            return str(r[c] if c is not None and c < len(r) and r[c] is not None else '').strip()
        name = get('name')
        if not name or field(name) == 'name' or re.match(r'^(合计|总计|小计|总价|客户签|结算状态|开单人|打印人|其他费用|优惠金额)', name):
            continue
        qty, price, amount = [number(get(k)) for k in ['qty', 'price', 'amount']]
        if all(x is None for x in [qty, price, amount]):
            if len(name) < 30:
                category = name
                room = '厨房' if '厨房' in name else name if re.fullmatch(r'.*(?:客厅|主卧|次卧|阳台|餐厅|儿童房|书房|玄关|卫生间)', name) else ''
            continue
        row = {k: get(k) for k in ALIASES}
        row.update(name=name, qty=qty, price=price, amount=amount,
                   room=get('room') if get('room') and number(get('room')) is None else room, category=get('category') or category)
        if amount is None and qty is not None and price is not None:
            row['amount'] = round(qty * price, 2)
            row['amountDerived'] = True
        if price is None and qty and amount is not None:
            row['price'] = amount / qty
            row['priceDerived'] = True
        if qty is None and price and amount is not None:
            row['qty'] = amount / price
            row['qtyDerived'] = True
        result.append(infer_material(row))
    return result, ''

def metadata(text, factory_mode=False):
    def match(pattern):
        m = re.search(pattern, text, re.I)
        return next((group.strip() for group in m.groups() if group is not None), '') if m else ''
    date = match(r'(20\d{2}[-/.年]\d{1,2}[-/.月]\d{1,2})')
    if date:
        parts = re.split(r'[-/.年月]', date)
        date = f'{parts[0]}-{int(parts[1]):02}-{int(parts[2]):02}'
    total_value = match(r'(?:(?:合计金额|总计金额|应付金额|实付金额)(?:\s*[（(]小写[）)])?)\s*[:：]?\s*([\d,.]+)')
    if not total_value and factory_mode:
        # Generic subtotal labels are easy to confuse with page counters (for
        # example, OCR can read “页小计 08” from “第1/1页”). Require currency.
        total_value = match(r'(?:合计|总计|页小计|小计)\s*[:：]?\s*([\d,.]+)\s*(?:元|￥|¥)')
    declared_total = number(total_value)
    if declared_total is not None and declared_total <= 0:
        declared_total = None
    return {'customer': match(r'(?:客户名称|客户姓名|客户|商户名称)[ \u3000]*[:：][ \u3000]*([^\n\t]+)'),
            'supplier': next((line.split('\t')[0].strip() for line in text.splitlines()[:12] if re.search(r'五金|板材|运营中心|销售中心' + (r'|新材' if factory_mode else ''), line) and len(line) < (65 if factory_mode else 55)), ''),
            'billNo': match(r'(?:订单号|单据编号|单号)\s*[:：]?\s*([A-Z\d-]{6,})'), 'date': date,
            'otherFee': number(match(r'其他费用\s*[:：]?\s*([\d.]+)')) or 0,
            'discount': number(match(r'(?:优惠金额|折让|优惠)\s*[:：]?\s*([\d.]+)')) or 0,
            'declaredTotal': declared_total}

_engine = None
_lock = threading.Lock()

def ocr_image(data, rotation=0, allow_headerless=False):
    import numpy as np
    from PIL import Image
    from rapidocr_onnxruntime import RapidOCR
    global _engine
    image = Image.open(io.BytesIO(data)).convert('RGB')
    if rotation:
        image = image.rotate(-rotation, expand=True)
    if image.width < 1200:
        scale = 1600 / image.width
        image = image.resize((1600, round(image.height * scale)), Image.Resampling.LANCZOS)
    with _lock:
        if _engine is None:
            _engine = RapidOCR(det_limit_side_len=2400, intra_op_num_threads=2, inter_op_num_threads=2)
        blocks, _ = _engine(np.array(image))
    blocks = blocks or []
    items = []
    for box, text, confidence in blocks:
        xs, ys = [p[0] for p in box], [p[1] for p in box]
        items.append({'x': (min(xs) + max(xs)) / 2, 'y': (min(ys) + max(ys)) / 2,
                      'h': max(ys) - min(ys), 'text': text, 'confidence': float(confidence)})
    items.sort(key=lambda b: b['y'])
    bands = []
    for b in items:
        band = next((g for g in reversed(bands[-3:]) if abs(g[0]['y'] - b['y']) <= max(9, b['h'] * .65)), None)
        if band is None:
            bands.append([b])
        else:
            band.append(b)
    for g in bands:
        g.sort(key=lambda b: b['x'])
    text = '\n'.join('\t'.join(b['text'] for b in g) for g in bands)
    headers, header_y = [], 0
    for band in bands:
        hits = [(field(b['text']), b['x']) for b in band if field(b['text'])]
        if any(k == 'name' for k, _ in hits) and len(hits) > len(headers):
            headers, header_y = hits, band[0]['y']
    grid = []
    if headers:
        headers.sort(key=lambda p: p[1])
        grid.append([ALIASES[k][0] for k, _ in headers])
        # Horizontal rules keep wrapped product names and numeric cells in their physical row.
        import cv2
        gray = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2GRAY)
        binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 21, 12)
        horizontal = cv2.morphologyEx(binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(30, image.width // 5), 1)))
        ys = np.where((horizontal > 0).sum(axis=1) > image.width * .45)[0]
        rules = []
        for y in ys:
            if not rules or y - rules[-1] > 3:
                rules.append(int(y))
        if len(rules) > 4 and any(y > header_y for y in rules):
            after = [y for y in rules if y > header_y]
            row_groups = [[] for _ in range(len(after) - 1)]
            for b in items:
                for i in range(len(after)-1):
                    if after[i] < b['y'] < after[i+1]:
                        row_groups[i].append(b)
                        break
        else:
            row_groups = [band for band in bands if band[0]['y'] > header_y + 8]
        for band in row_groups:
            cells = [''] * len(headers)
            for b in sorted(band, key=lambda b: (round(b['y']/8), b['x'])):
                col = min(range(len(headers)), key=lambda n: abs(headers[n][1] - b['x']))
                cells[col] = (cells[col] + ' ' + b['text']).strip()
            grid.append(cells)
    # A wrapped specification may fall below the detected horizontal rule.
    # Join only unnamed continuation cells, never rows with their own product.
    if headers and grid:
        name_col = next((i for i, (k, _) in enumerate(headers) if k == 'name'), None)
        if name_col is not None:
            merged = [grid[0]]
            for cells in grid[1:]:
                footer = any(re.search(r'结算状态|合计金额|其他费用|优惠金额', c) for c in cells)
                if not cells[name_col] and not footer and len(merged) > 1 and merged[-1][name_col]:
                    for i, cell in enumerate(cells):
                        if cell:
                            merged[-1][i] = (merged[-1][i] + ' ' + cell).strip()
                else:
                    merged.append(cells)
            grid = merged
    if not headers and allow_headerless:
        # Continuation pages commonly omit column labels. Infer the ruled
        # 8-column purchase-slip layout and leave its OCR values editable.
        import cv2
        gray = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2GRAY)
        binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 21, 12)
        horizontal = cv2.morphologyEx(binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(30, image.width // 5), 1)))
        # Continuation bills can contain only two or three item rows, so use a
        # shorter vertical-line kernel than a full-page form needs.
        vertical = cv2.morphologyEx(binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(24, image.height // 12))))
        def line_centers(values, threshold, gap=3):
            coords = np.where(values > threshold)[0]
            groups = []
            for value in coords:
                if not groups or value - groups[-1][-1] > gap:
                    groups.append([int(value)])
                else:
                    groups[-1].append(int(value))
            return [int(sum(group) / len(group)) for group in groups]
        x_rules = line_centers((vertical > 0).sum(axis=0), image.height * .12, max(3, image.width // 800))
        y_rules = line_centers((horizontal > 0).sum(axis=1), image.width * .30, 3)
        intervals = len(x_rules) - 1
        row_bands = list(zip(y_rules, y_rules[1:]))
        serial_rows = sum(any(re.fullmatch(r'\d{1,3}', b['text'].strip()) for b in items
                              if y0 < b['y'] < y1 and x_rules[0] <= b['x'] < x_rules[1])
                          for y0, y1 in row_bands) if intervals >= 7 else 0
        has_serial = serial_rows >= 2
        labels = ['序号', '品名', '规格', '单位', '数量', '单价', '金额', '备注'] if intervals == 8 or intervals == 7 and has_serial else ['品名', '规格', '单位', '数量', '单价', '金额', '备注']
        if intervals in (7, 8) and len(row_bands) >= 2:
            grid = [labels]
            for y0, y1 in row_bands:
                cells = [''] * len(labels)
                for block in items:
                    if not y0 < block['y'] < y1:
                        continue
                    col = next((i for i in range(len(x_rules) - 1) if x_rules[i] <= block['x'] < x_rules[i + 1]), None)
                    if col is not None:
                        cells[col] = (cells[col] + ' ' + block['text']).strip()
                if any(cells):
                    grid.append(cells)
    rows, warning = table_rows(grid)
    if allow_headerless and not rows and rotation == 0:
        # Factory suppliers sometimes send sideways phone photos; retry quarter
        # turns only when the normal orientation finds no usable rows.
        for angle in (90, 270, 180):
            candidate = ocr_image(data, angle, True)
            if candidate[0]:
                return candidate
    brands = ['福人', '千年舟', '兔宝宝', '大亚圣象', '禾香'] + (['华木'] if allow_headerless else [])
    supplier_brand = next((x for x in brands if x in '\n'.join(text.splitlines()[:3])), '')
    if supplier_brand:
        for row in rows:
            if not row.get('brand') and not re.search(r'打包|封边|包装|运费', row['name']):
                row['brand'] = supplier_brand
    return rows, text, grid, warning

def extract(data, filename, rotation=0, factory_mode=False):
    ext = Path(filename).suffix.lower()
    sheets, text, warnings = [], '', []
    if ext in ['.xlsx', '.xlsm']:
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for ws in book:
            grid = [[v if v is not None else '' for v in r] for r in ws.iter_rows(values_only=True)]
            # Expand actual merged ranges, not unrelated blank room cells.
            for merged in ws.merged_cells.ranges:
                if merged.min_col != merged.max_col:
                    continue
                value = grid[merged.min_row - 1][merged.min_col - 1]
                for ri in range(merged.min_row - 1, merged.max_row):
                    for ci in range(merged.min_col - 1, merged.max_col):
                        grid[ri][ci] = value
            sheets.append((ws.title, grid))
        book.close()
    elif ext == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sheets = [(ws.name, [ws.row_values(i) for i in range(ws.nrows)]) for ws in book.sheets()]
    elif ext in ['.csv', '.tsv', '.txt']:
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('gb18030')
        sep = '\t' if '\t' in text else ','
        sheets = [(filename, list(csv.reader(io.StringIO(text), delimiter=sep)))]
    elif ext == '.pdf':
        import fitz
        book = fitz.open(stream=data, filetype='pdf')
        if len(book) > 20:
            raise ValueError('请将 PDF 分为每份不超过20页的单据。')
        all_rows = []
        for page in book:
            text += page.get_text() + '\n'
            tables = page.find_tables()
            if tables.tables:
                for table in tables.tables:
                    grid = table.extract()
                    rows, warn = table_rows(grid)
                    all_rows.extend(rows)
                    sheets.append((f'第{page.number+1}页', grid))
            else:
                rows, txt, grid, warn = ocr_image(page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png'), rotation, factory_mode)
                all_rows.extend(rows)
                sheets.append((f'第{page.number+1}页', grid))
                text += txt + '\n'
        book.close()
        if all_rows:
            return {'rows': all_rows, 'text': text, 'sheets': [{'name': n, 'grid': g} for n, g in sheets], 'meta': metadata(text, factory_mode), 'warnings': warnings}
    else:
        rows, text, grid, warning = ocr_image(data, rotation, factory_mode)
        return {'rows': rows, 'text': text, 'sheets': [{'name': filename, 'grid': grid}], 'meta': metadata(text, factory_mode), 'warnings': [warning] if warning else []}
    result = []
    for name, grid in sheets:
        rows, warning = table_rows(grid)
        result.extend({**r, 'sheet': name} for r in rows)
        if warning:
            warnings.append(name + '：' + warning)
    if not text:
        text = '\n'.join('\t'.join(str(v) for v in r) for _, grid in sheets for r in grid)
    return {'rows': result, 'text': text, 'sheets': [{'name': n, 'grid': g[:1500]} for n, g in sheets], 'meta': metadata(text, factory_mode), 'warnings': warnings}

class Handler(BaseHTTPRequestHandler):
    def reply(self, status, data):
        body = json.dumps(data, ensure_ascii=False, allow_nan=False, default=lambda v: v.isoformat() if isinstance(v, (date, datetime)) else str(v)).encode('utf-8')
        self.send_response(status)
        origin = self.headers.get('Origin', '')
        if origin in ['http://127.0.0.1:8765', 'http://localhost:8765']:
            self.send_header('Access-Control-Allow-Origin', origin)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self):
        self.send_response(204)
        origin = self.headers.get('Origin', '')
        if origin in ['http://127.0.0.1:8765', 'http://localhost:8765']:
            self.send_header('Access-Control-Allow-Origin', origin)
        self.send_header('Access-Control-Allow-Methods', 'POST, GET, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def do_GET(self):
        self.reply(200, {'status': 'ready', 'service': 'anju-local-document-import'})

    def do_POST(self):
        try:
            if self.headers.get('Origin', '') not in ['', 'http://127.0.0.1:8765', 'http://localhost:8765']:
                return self.reply(403, {'error': '仅允许安居本机预览调用。'})
            if urlsplit(self.path).path != '/extract':
                return self.reply(404, {'error': '接口不存在'})
            size = int(self.headers.get('Content-Length', 0))
            if size <= 0 or size > 30 * 1024 * 1024:
                raise ValueError('请选择小于30MB的单据。')
            q = parse_qs(urlsplit(self.path).query)
            data = self.rfile.read(size)
            self.reply(200, extract(data, q.get('name', ['document.png'])[0], int(q.get('rotation', ['0'])[0]), q.get('factory', ['0'])[0] == '1'))
        except Exception as exc:
            self.reply(400, {'error': str(exc)})

if __name__ == '__main__':
    ThreadingHTTPServer(('127.0.0.1', 8766), Handler).serve_forever()
